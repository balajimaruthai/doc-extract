"""
DOCX text extraction using python-docx.
Preserves paragraph structure and table content.
"""
import logging
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


@dataclass
class DocxExtractionResult:
    document_path: str
    full_text: str
    paragraph_count: int
    table_count: int
    error: Optional[str] = None


def extract_docx(file_path: Path) -> DocxExtractionResult:
    """
    Extract text from a DOCX file.
    Captures: paragraphs, headings, table cells.
    """
    try:
        from docx import Document
    except ImportError:
        return DocxExtractionResult(
            document_path=str(file_path),
            full_text="",
            paragraph_count=0,
            table_count=0,
            error="python-docx is not installed",
        )

    try:
        doc = Document(str(file_path))
    except Exception as e:
        logger.error(f"Failed to open DOCX {file_path}: {e}")
        return DocxExtractionResult(
            document_path=str(file_path),
            full_text="",
            paragraph_count=0,
            table_count=0,
            error=f"Cannot open DOCX file: {e}",
        )

    lines: list[str] = []
    para_count = 0
    table_count = 0

    # Extract paragraphs (preserves order)
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)
            para_count += 1

    # Extract table content
    for table in doc.tables:
        table_count += 1
        lines.append("")  # Visual separator
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                lines.append(" | ".join(row_cells))
        lines.append("")

    full_text = "\n".join(lines).strip()

    logger.info(
        f"DOCX extracted: {file_path.name} | "
        f"{para_count} paragraphs | {table_count} tables | "
        f"{len(full_text)} chars"
    )

    return DocxExtractionResult(
        document_path=str(file_path),
        full_text=full_text,
        paragraph_count=para_count,
        table_count=table_count,
    )
