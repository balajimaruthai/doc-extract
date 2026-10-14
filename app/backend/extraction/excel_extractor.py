"""
Excel and Spreadsheet text/data extraction using openpyxl, pandas, and csv.
Extracts tabular structure, sheet names, line items, and calculates financial sums.
Zero cloud dependencies — 100% offline.
"""
import csv
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Optional

logger = logging.getLogger(__name__)


@dataclass
class SpreadsheetSheet:
    sheet_name: str
    row_count: int
    column_count: int
    headers: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    calculated_totals: dict[str, float] = field(default_factory=dict)
    text_representation: str = ""


@dataclass
class ExcelExtractionResult:
    document_path: str
    sheets: list[SpreadsheetSheet] = field(default_factory=list)
    full_text: str = ""
    total_rows: int = 0
    total_sheets: int = 0
    grand_total: Optional[float] = None
    total_credits: Optional[float] = None
    total_debits: Optional[float] = None
    currency_symbol: str = "$"
    detected_financial_summary: dict[str, Any] = field(default_factory=dict)
    error: Optional[str] = None


def extract_spreadsheet(file_path: Path) -> ExcelExtractionResult:
    """
    Extract structured text and financial data from an Excel (.xlsx, .xls) or CSV file.
    """
    path = Path(file_path)
    suffix = path.suffix.lower().lstrip(".")

    if suffix == "csv":
        return _extract_csv(path)
    elif suffix in ("xlsx", "xlsm", "xltx"):
        return _extract_xlsx(path)
    elif suffix == "xls":
        return _extract_xls(path)
    else:
        return ExcelExtractionResult(
            document_path=str(path),
            error=f"Unsupported spreadsheet format: .{suffix}",
        )


def _extract_xlsx(file_path: Path) -> ExcelExtractionResult:
    """Extract from modern Excel (.xlsx)."""
    try:
        import openpyxl

        wb = openpyxl.load_workbook(str(file_path), data_only=True, read_only=True)
        sheets: list[SpreadsheetSheet] = []
        all_text_parts: list[str] = []
        overall_grand_total = 0.0
        total_debits = 0.0
        total_credits = 0.0
        has_grand_total = False

        for name in wb.sheetnames:
            ws = wb[name]
            raw_rows = []
            for r in ws.iter_rows(values_only=True):
                # Clean row values to string
                str_row = ["" if v is None else str(v).strip() for v in r]
                # Skip completely blank rows
                if any(str_row):
                    raw_rows.append(str_row)

            if not raw_rows:
                continue

            headers = [h for h in raw_rows[0]]
            data_rows = raw_rows[1:]

            sheet_res = _process_sheet_data(name, headers, data_rows)
            sheets.append(sheet_res)
            all_text_parts.append(sheet_res.text_representation)

            for col_name, total_val in sheet_res.calculated_totals.items():
                col_lower = col_name.lower()
                if any(k in col_lower for k in ["total", "amount", "net", "due", "balance", "price"]):
                    overall_grand_total += total_val
                    has_grand_total = True
                if "credit" in col_lower or "deposit" in col_lower:
                    total_credits += total_val
                if "debit" in col_lower or "withdrawal" in col_lower or "expense" in col_lower:
                    total_debits += total_val

        wb.close()
        full_text = "\n\n".join(all_text_parts)

        return ExcelExtractionResult(
            document_path=str(file_path),
            sheets=sheets,
            full_text=full_text,
            total_rows=sum(s.row_count for s in sheets),
            total_sheets=len(sheets),
            grand_total=round(overall_grand_total, 2) if has_grand_total else None,
            total_credits=round(total_credits, 2) if total_credits > 0 else None,
            total_debits=round(total_debits, 2) if total_debits > 0 else None,
        )

    except Exception as e:
        logger.warning(f"openpyxl failed on {file_path.name}: {e}. Trying pandas...")
        return _extract_via_pandas(file_path)


def _extract_xls(file_path: Path) -> ExcelExtractionResult:
    """Extract legacy Excel (.xls) using pandas."""
    return _extract_via_pandas(file_path)


def _extract_via_pandas(file_path: Path) -> ExcelExtractionResult:
    """Fallback extraction using pandas."""
    try:
        import pandas as pd

        excel_file = pd.ExcelFile(str(file_path))
        sheets: list[SpreadsheetSheet] = []
        all_text_parts: list[str] = []
        overall_grand_total = 0.0
        has_grand_total = False

        for sheet_name in excel_file.sheet_names:
            df = pd.read_excel(excel_file, sheet_name=sheet_name)
            df = df.dropna(how="all")
            headers = [str(c) for c in df.columns]
            data_rows = [[str(v) if pd.notna(v) else "" for v in row] for row in df.values]

            sheet_res = _process_sheet_data(sheet_name, headers, data_rows)
            sheets.append(sheet_res)
            all_text_parts.append(sheet_res.text_representation)

            for col_name, total_val in sheet_res.calculated_totals.items():
                if any(k in col_name.lower() for k in ["total", "amount", "net", "due", "balance"]):
                    overall_grand_total += total_val
                    has_grand_total = True

        return ExcelExtractionResult(
            document_path=str(file_path),
            sheets=sheets,
            full_text="\n\n".join(all_text_parts),
            total_rows=sum(s.row_count for s in sheets),
            total_sheets=len(sheets),
            grand_total=round(overall_grand_total, 2) if has_grand_total else None,
        )
    except Exception as err:
        logger.error(f"Pandas extraction failed on {file_path}: {err}")
        return ExcelExtractionResult(
            document_path=str(file_path),
            error=f"Spreadsheet read error: {err}",
        )


def _extract_csv(file_path: Path) -> ExcelExtractionResult:
    """Extract from CSV file."""
    try:
        lines: list[list[str]] = []
        for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                with open(file_path, "r", encoding=encoding, errors="replace") as f:
                    # Detect delimiter
                    sample = f.read(4096)
                    f.seek(0)
                    dialect = csv.excel
                    try:
                        sniffer = csv.Sniffer()
                        dialect = sniffer.sniff(sample)
                    except Exception:
                        pass
                    reader = csv.reader(f, dialect=dialect)
                    lines = [row for row in reader if any(cell.strip() for cell in row)]
                break
            except Exception:
                continue

        if not lines:
            return ExcelExtractionResult(document_path=str(file_path), full_text="")

        headers = lines[0]
        data_rows = lines[1:]
        sheet_res = _process_sheet_data("CSV_Data", headers, data_rows)

        has_grand_total = False
        overall_grand_total = 0.0
        for col_name, total_val in sheet_res.calculated_totals.items():
            if any(k in col_name.lower() for k in ["total", "amount", "net", "due", "balance", "price"]):
                overall_grand_total += total_val
                has_grand_total = True

        return ExcelExtractionResult(
            document_path=str(file_path),
            sheets=[sheet_res],
            full_text=sheet_res.text_representation,
            total_rows=sheet_res.row_count,
            total_sheets=1,
            grand_total=round(overall_grand_total, 2) if has_grand_total else None,
        )
    except Exception as err:
        return ExcelExtractionResult(
            document_path=str(file_path),
            error=f"CSV read error: {err}",
        )


def _process_sheet_data(sheet_name: str, headers: list[str], rows: list[list[str]]) -> SpreadsheetSheet:
    """Analyze headers and rows, compute column totals, build markdown text."""
    calculated_totals: dict[str, float] = {}
    numeric_col_indices: dict[int, str] = {}

    # Identify numeric columns (Amount, Total, Gross, Price, Debit, Credit, etc.)
    for idx, h in enumerate(headers):
        h_clean = h.strip()
        if any(k in h_clean.lower() for k in [
            "amount", "total", "sum", "price", "cost", "gross", "net",
            "debit", "credit", "balance", "fee", "tax", "charge", "salary", "pay"
        ]):
            numeric_col_indices[idx] = h_clean

    # Sum columns
    for idx, col_name in numeric_col_indices.items():
        col_sum = 0.0
        count = 0
        for r in rows:
            if idx < len(r):
                val_str = r[idx]
                clean_num = re.sub(r"[^\d.-]", "", val_str.replace(",", ""))
                try:
                    if clean_num:
                        col_sum += float(clean_num)
                        count += 1
                except ValueError:
                    pass
        if count > 0:
            calculated_totals[col_name] = round(col_sum, 2)

    # Build clean formatted text representation
    text_lines = [f"=== SHEET: {sheet_name} ({len(rows)} rows) ==="]

    # Add header
    if headers:
        text_lines.append(" | ".join(headers))
        text_lines.append("-|-".join(["---"] * len(headers)))

    # Add up to first 100 rows for text/LLM processing
    max_preview_rows = min(100, len(rows))
    for r in rows[:max_preview_rows]:
        text_lines.append(" | ".join(r))

    if len(rows) > max_preview_rows:
        text_lines.append(f"... [{len(rows) - max_preview_rows} more rows in sheet] ...")

    # Add calculated column totals
    if calculated_totals:
        text_lines.append("\n--- CALCULATED TOTALS & SUMMARY ---")
        for col_name, total_val in calculated_totals.items():
            text_lines.append(f"Total {col_name}: ${total_val:,.2f}")

    sheet_text = "\n".join(text_lines)

    return SpreadsheetSheet(
        sheet_name=sheet_name,
        row_count=len(rows),
        column_count=len(headers),
        headers=headers,
        rows=rows,
        calculated_totals=calculated_totals,
        text_representation=sheet_text,
    )
